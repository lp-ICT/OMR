"""Local-only OMR reviewer for the school's F5ICT 60-question answer sheet."""
import csv
import io
import json
import os
import threading
import urllib.parse
import webbrowser
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import cv2
import fitz
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from PIL import Image

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
DATA.mkdir(exist_ok=True)
STATE_FILE = DATA / "session.json"
WIDTH, HEIGHT = 990, 1400
LETTERS = "ABCD"
MAX_UPLOAD = 250 * 1024 * 1024


def fresh():
    return {"key": [""] * 60, "students": [], "roster": [], "audit": [], "source": ""}


def load():
    try:
        obj = json.loads(STATE_FILE.read_text(encoding="utf-8"))
        return obj if isinstance(obj, dict) and "students" in obj else fresh()
    except (OSError, ValueError):
        return fresh()


state = load()
lock = threading.RLock()


def save():
    temporary = STATE_FILE.with_suffix(".tmp")
    temporary.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
    temporary.replace(STATE_FILE)


def rect_dark(gray, x, y, rx=8, ry=3):
    crop = gray[max(0, int(y-ry)):int(y+ry+1), max(0, int(x-rx)):int(x+rx+1)]
    return float(1 - np.mean(crop) / 255) if crop.size else 0.0


def answer_positions(q):
    col = 0 if q < 30 else 1
    within = (q % 30)
    group, row = divmod(within, 10)
    y = round(510 + group * 280 + row * 25.65)
    return [(x, y) for x in ([444, 486, 527, 568] if col == 0 else [730, 770, 811, 851])]


def classify(scores):
    order = sorted(range(len(scores)), key=lambda k: scores[k], reverse=True)
    strongest, second = scores[order[0]], scores[order[1]]
    candidates = [i for i, v in enumerate(scores) if v >= max(0.30, strongest * 0.68)]
    if strongest < 0.27:
        return "", "漏答"
    if len(candidates) > 1:
        return "?", "多選／擦改"
    if strongest < 0.38 or strongest-second < 0.15:
        return LETTERS[order[0]], "低信心"
    return LETTERS[order[0]], ""


def read_sheet(gray):
    items = []
    offsets = []
    for group in range(6):
        ranked = []
        for dy in range(-8, 9):
            contrast = 0.0
            for q in range(group * 10, group * 10 + 10):
                values = sorted((rect_dark(gray, x, y+dy) for x, y in answer_positions(q)), reverse=True)
                contrast += values[0] - values[1]
            ranked.append((contrast, dy))
        offsets.append(max(ranked)[1])
    for q in range(60):
        # Fit each block of ten answers to its printed grid before sampling.
        dy = offsets[q // 10]
        scores = [round(rect_dark(gray, x, y+dy), 3) for x, y in answer_positions(q)]
        ans, flag = classify(scores)
        items.append({"auto": ans, "answer": ans, "flag": flag, "scores": scores, "reviewed": False})
    digits = []
    identity_flags = []
    for x in (704, 739):
        scores = [rect_dark(gray, x, round(247 + n * 18.7), 3, 4) for n in range(10)]
        order = sorted(range(10), key=lambda n: scores[n], reverse=True)
        if scores[order[0]] < 0.29 or scores[order[0]]-scores[order[1]] < 0.10:
            digits.append("?")
            identity_flags.append("學號格不清楚")
        else:
            digits.append(str(order[0]))
    return items, "".join(digits), identity_flags


def find_roster(student_no, roster):
    matches = [r for r in roster if r["學號"].lstrip("0") == student_no.lstrip("0") and student_no and "?" not in student_no]
    return matches[0] if len(matches) == 1 else None


def refresh_roster():
    for s in state["students"]:
        match = find_roster(s["number"], state["roster"])
        if match and not s.get("identity_edited"):
            s["class"] = match["班別"]
            s["name"] = match["姓名"]


def make_page(page, rotation):
    scale = WIDTH / page.rect.width
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    if rotation:
        img = img.rotate(-rotation, expand=True)
    img = img.resize((WIDTH, HEIGHT), Image.Resampling.LANCZOS)
    return img


def csv_roster(raw):
    text = None
    for encoding in ("utf-8-sig", "big5", "utf-16"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeError:
            pass
    if text is None:
        raise ValueError("名冊編碼須為 UTF-8、Big5 或 UTF-16")
    rows = list(csv.DictReader(io.StringIO(text)))
    if not all(k in (rows[0] if rows else {}) for k in ("班別", "學號", "姓名")):
        raise ValueError("CSV 欄名必須包含：班別,學號,姓名")
    return [{k: str(row.get(k, "")).strip() for k in ("班別", "學號", "姓名")} for row in rows]


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def prepare_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "學生總表"
    ws.append(["頁數", "班別", "學號", "姓名", "得分", "待覆核題數", "身份提醒"] + [f"第{i}題" for i in range(1, 61)])
    key = state["key"]
    ws.append(["標準答案", "", "", "", "", "", ""] + key)
    item = wb.create_sheet("逐題分析")
    item.append(["題目", "標準答案", "答對", "答錯", "漏答", "多選／不確定", "答對率", "A", "B", "C", "D"])
    detail = wb.create_sheet("辨識及覆核紀錄")
    detail.append(["頁數", "學號", "題目", "原辨識", "最後答案", "自動提示", "人工覆核", "A密度", "B密度", "C密度", "D密度"])
    audit = wb.create_sheet("修改紀錄")
    audit.append(["時間", "頁數", "項目", "修改前", "修改後"])
    for a in state["audit"]:
        audit.append([a.get(k, "") for k in ("time", "page", "field", "old", "new")])
    for s in state["students"]:
        answers = [a["answer"] for a in s["answers"]]
        score = sum(bool(key[i]) and answers[i] == key[i] for i in range(60))
        flags = sum(bool(key[i]) and bool(a["flag"]) and not a["reviewed"]
                    for i, a in enumerate(s["answers"]))
        ws.append([s["page"], safe_cell(s["class"]), safe_cell(s["number"]), safe_cell(s["name"]), score, flags,
                   "、".join(s["identity_flags"])] + answers)
        for i, a in enumerate(s["answers"], 1):
            detail.append([s["page"], safe_cell(s["number"]), i, a["auto"], a["answer"], a["flag"],
                           "是" if a["reviewed"] else "否", *a["scores"]])
    for i in range(60):
        answers = [s["answers"][i]["answer"] for s in state["students"]]
        right = sum(bool(key[i]) and a == key[i] for a in answers)
        blank = answers.count("")
        uncertain = answers.count("?")
        item.append([i+1, key[i], right, len(answers)-right-blank-uncertain,
                     blank, uncertain, right/len(answers) if answers and key[i] else None,
                     *[answers.count(c) for c in LETTERS]])
        item.cell(i+2, 7).number_format = "0.0%"
    for sheet in wb:
        sheet.freeze_panes = "H3" if sheet == ws else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 26
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="0A518A")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(23, max(10, max(len(str(c.value or "")) for c in list(col)[:100])+2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        pass

    def output(self, code, data, content_type="application/json", filename=None):
        if isinstance(data, (dict, list)):
            data = json.dumps(data, ensure_ascii=False).encode("utf-8")
        if isinstance(data, str):
            data = data.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        if filename:
            self.send_header("Content-Disposition", f"attachment; filename={filename}")
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        path = urllib.parse.urlsplit(self.path).path
        try:
            if path == "/":
                return self.output(200, (ROOT / "index.html").read_bytes(), "text/html; charset=utf-8")
            if path == "/logo.png":
                return self.output(200, (ROOT / "logo.png").read_bytes(), "image/png")
            if path == "/api/state":
                with lock:
                    return self.output(200, state)
            if path.startswith("/api/page/"):
                n = int(path.rsplit("/", 1)[1])
                if not (1 <= n <= len(state["students"])):
                    raise ValueError("頁數不存在")
                return self.output(200, (DATA / f"page_{n}.jpg").read_bytes(), "image/jpeg")
            if path == "/api/export":
                with lock:
                    return self.output(200, prepare_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "OMR_results.xlsx")
            return self.output(404, {"error": "找不到頁面"})
        except Exception as e:
            self.output(400, {"error": str(e)})

    def do_POST(self):
        path = urllib.parse.urlsplit(self.path).path
        try:
            size = int(self.headers.get("Content-Length", "0"))
            if size < 1 or size > MAX_UPLOAD:
                raise ValueError("檔案大小須少於 250 MB")
            raw = self.rfile.read(size)
            with lock:
                if path == "/api/upload":
                    params = urllib.parse.parse_qs(urllib.parse.urlsplit(self.path).query)
                    rotation = int(params.get("rotation", [0])[0])
                    if rotation not in (0, 90, 180, 270):
                        raise ValueError("旋轉角度錯誤")
                    doc = fitz.open(stream=raw, filetype="pdf")
                    if not 1 <= len(doc) <= 300:
                        raise ValueError("PDF 頁數須為 1 至 300")
                    first_key = params.get("firstKey", ["0"])[0] == "1"
                    new_students = []
                    auto_key = None
                    for idx, page in enumerate(doc):
                        img = make_page(page, rotation)
                        gray = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2GRAY)
                        answers, number, identity_flags = read_sheet(gray)
                        if idx == 0 and first_key:
                            auto_key = [a["answer"] if a["answer"] in LETTERS else "" for a in answers]
                            continue
                        page_number = len(new_students)+1
                        img.save(DATA / f"page_{page_number}.jpg", quality=86)
                        matched = find_roster(number, state["roster"])
                        new_students.append({"page": page_number, "source_page": idx+1,
                            "number": number, "class": matched["班別"] if matched else "",
                            "name": matched["姓名"] if matched else "", "identity_flags": identity_flags,
                            "identity_edited": False, "answers": answers})
                    if not new_students:
                        raise ValueError("PDF 沒有學生答卷")
                    state["students"] = new_students
                    state["source"] = "已載入 PDF（" + str(len(doc)) + " 頁）"
                    state["audit"] = []
                    if auto_key:
                        state["key"] = auto_key
                    save()
                    return self.output(200, {"count": len(new_students)})
                if path == "/api/roster":
                    state["roster"] = csv_roster(raw)
                    refresh_roster()
                    save()
                    return self.output(200, {"count": len(state["roster"])})
                obj = json.loads(raw)
                if path == "/api/key":
                    key = obj.get("key")
                    if not isinstance(key, list) or len(key) != 60 or any(x not in ("", *LETTERS) for x in key):
                        raise ValueError("標準答案必須為 60 題，每題 A–D 或留空")
                    state["key"] = key
                    save()
                    return self.output(200, {"ok": True})
                if path == "/api/update":
                    page, field = int(obj["page"]), obj["field"]
                    s = state["students"][page-1]
                    value = str(obj.get("value", "")).strip()
                    if field == "answer":
                        q = int(obj["question"])-1
                        if not 0 <= q < 60 or value not in ("", "?", *LETTERS):
                            raise ValueError("答案無效")
                        a = s["answers"][q]
                        old = a["answer"]
                        a["answer"] = value
                        a["reviewed"] = True
                        label = f"第{q+1}題"
                    elif field in ("class", "number", "name"):
                        if len(value) > 40:
                            raise ValueError("身份欄位過長")
                        old = s[field]
                        s[field] = value
                        s["identity_edited"] = True
                        s["identity_flags"] = []
                        label = field
                    else:
                        raise ValueError("未知欄位")
                    if old != value:
                        state["audit"].append({"time": datetime.now().isoformat(timespec="seconds"),
                             "page": page, "field": label, "old": old, "new": value})
                    save()
                    return self.output(200, {"ok": True})
            return self.output(404, {"error": "找不到操作"})
        except Exception as e:
            self.output(400, {"error": str(e)})


if __name__ == "__main__":
    url = "http://127.0.0.1:8765/"
    print("OMR 已啟動：" + url + "；結束請按 Ctrl+C")
    threading.Timer(1.2, lambda: webbrowser.open(url)).start()
    ThreadingHTTPServer(("127.0.0.1", 8765), Handler).serve_forever()
